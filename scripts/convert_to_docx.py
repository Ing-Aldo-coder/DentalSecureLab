import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def markdown_to_docx(md_path, docx_path, base_image_dir):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(6)

    with open(md_path, 'r', encoding='utf-8', errors='replace') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Determine cols
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        for r_idx, row_data in enumerate(table_rows):
            row = table.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell.text = text.strip()
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

                # Style cell text
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if is_header:
                    set_cell_background(cell, "0F172A") # Slate 900
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    for run in p.runs:
                        run.font.size = Pt(9.0)
                        run.font.color.rgb = RGBColor(30, 41, 59)

        doc.add_paragraph() # space after table
        in_table = False
        table_rows = []

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "1E1E1E") # Dark IDE background
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        full_code = "".join(code_lines).rstrip()
        run = p.add_run(full_code)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(220, 220, 220)

        doc.add_paragraph() # spacing
        in_code_block = False
        code_lines = []

    for line in lines:
        raw_line = line
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        # Handle tables
        if '|' in stripped and not stripped.startswith('#') and not stripped.startswith('!['):
            # Check separator row
            if re.match(r'^[\|\s\:\-]+$', stripped):
                continue
            parts = [c.strip() for c in stripped.strip('|').split('|')]
            if len(parts) >= 2:
                in_table = True
                table_rows.append(parts)
                continue
        elif in_table:
            flush_table()

        # Handle horizontal rules
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("_________________________________________________________________________________")
            run.font.color.rgb = RGBColor(203, 213, 225)
            continue

        # Handle empty lines
        if not stripped:
            continue

        # Handle Headings
        if stripped.startswith('# '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[2:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
            continue
        elif stripped.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[3:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199) # Sky 600
            continue
        elif stripped.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[4:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
            continue
        elif stripped.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[5:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
            continue

        # Handle Images ![alt](src)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_rel_path = img_match.group(2)
            
            # Resolve image full path
            if os.path.isabs(img_rel_path):
                img_full_path = img_rel_path
            else:
                img_full_path = os.path.normpath(os.path.join(base_image_dir, img_rel_path))
            
            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.add_run().add_picture(img_full_path, width=Inches(5.8))
                
                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(f"Evidencia: {alt_text}")
                run_cap.font.size = Pt(8.5)
                run_cap.font.italic = True
                run_cap.font.color.rgb = RGBColor(100, 116, 139)
            else:
                print(f"Warning: image not found at {img_full_path}")
            continue

        # Handle list items
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            clean_text = stripped[2:].strip()
            format_inline_markdown(p, clean_text)
            continue
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            clean_text = re.sub(r'^\d+\.\s', '', stripped).strip()
            format_inline_markdown(p, clean_text)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        format_inline_markdown(p, stripped)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f"Document successfully created: {docx_path}")

def format_inline_markdown(paragraph, text):
    # Splits text by bold **...** and code `...`
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(194, 65, 12) # Dark orange/red code
        else:
            run = paragraph.add_run(token)
            run.font.color.rgb = RGBColor(51, 65, 85)

if __name__ == '__main__':
    base_dir = r'C:\Users\Zephyrus\Desktop\practica\DentalSecureLab\docs'
    target_dir = r'C:\Users\Zephyrus\Desktop\practica'

    docs_to_convert = [
        ('REPORTE_TECNICO_AUDITORIA.md', 'REPORTE_TECNICO_AUDITORIA.docx'),
        ('manual_usuario_seguro.md', 'MANUAL_USUARIO_SEGURO.docx'),
        ('politicas_y_matrices_seguridad.md', 'POLITICAS_Y_MATRICES_SEGURIDAD.docx'),
        ('arquitectura.md', 'ARQUITECTURA_SISTEMA.docx'),
        ('sop_antiphishing.md', 'SOP_ANTIPHISHING.docx'),
        ('network_hardening.md', 'NETWORK_HARDENING.docx')
    ]

    for src, dst in docs_to_convert:
        src_path = os.path.join(base_dir, src)
        dst_path = os.path.join(target_dir, dst)
        if os.path.exists(src_path):
            print(f"Converting {src} -> {dst}...")
            try:
                markdown_to_docx(src_path, dst_path, base_dir)
            except PermissionError:
                print(f"  [AVISO] {dst} no pudo sobrescribirse porque está abierto en Microsoft Word. Ciérralo si deseas regenerarlo.")
            except Exception as e:
                print(f"  [ERROR] en {src}: {e}")
        else:
            print(f"Source not found: {src_path}")
